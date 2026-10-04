from pathlib import Path
import csv
import io

import pymupdf
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
import pandas as pd
from PIL import Image
import pytesseract


CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150


# ============================================================
# PDF
# ============================================================

def extract_pdf(file_path: str) -> list[dict]:
    """Extract text from a PDF. Uses OCR when a page has no text."""
    pages = []

    with pymupdf.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):

            text = page.get_text("text").strip()

            # Normal PDF with selectable text
            if text:
                pages.append({
                    "page": page_number,
                    "text": text,
                    "source_type": "pdf_text",
                })
                continue

            # Scanned/image-only PDF → OCR
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))

            image = Image.open(
                io.BytesIO(pixmap.tobytes("png"))
            )

            ocr_text = pytesseract.image_to_string(image).strip()

            if ocr_text:
                pages.append({
                    "page": page_number,
                    "text": ocr_text,
                    "source_type": "pdf_ocr",
                })

    return pages


# ============================================================
# TXT
# ============================================================

def extract_txt(file_path: str) -> list[dict]:
    """Extract text from a plain text file."""
    path = Path(file_path)

    text = path.read_text(
        encoding="utf-8-sig",
        errors="replace"
    ).strip()

    if not text:
        return []

    return [{
        "page": None,
        "text": text,
        "source_type": "text",
    }]


# ============================================================
# DOCX
# ============================================================

def extract_docx(file_path: str) -> list[dict]:
    """Extract paragraphs and tables from a DOCX file."""
    document = Document(file_path)

    parts = []

    # Paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(
                cell for cell in cells if cell
            )

            if row_text:
                parts.append(row_text)

    text = "\n".join(parts).strip()

    if not text:
        return []

    return [{
        "page": None,
        "text": text,
        "source_type": "docx",
    }]


# ============================================================
# PPTX
# ============================================================

def extract_pptx(file_path: str) -> list[dict]:
    """Extract text from PowerPoint slides."""
    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        parts = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    parts.append(text)

        slide_text = "\n".join(parts).strip()

        if slide_text:
            slides.append({
                "page": slide_number,
                "text": slide_text,
                "source_type": "pptx",
            })

    return slides


# ============================================================
# XLSX / XLS
# ============================================================

def extract_excel(file_path: str) -> list[dict]:
    """Extract spreadsheet data sheet-by-sheet."""
    path = Path(file_path)

    sheets = []

    try:
        excel_file = pd.ExcelFile(file_path)

        for sheet_name in excel_file.sheet_names:

            dataframe = pd.read_excel(
                file_path,
                sheet_name=sheet_name,
                header=None,
            )

            dataframe = dataframe.fillna("")

            rows = []

            for row in dataframe.values.tolist():

                values = [
                    str(value).strip()
                    for value in row
                    if str(value).strip()
                ]

                if values:
                    rows.append(" | ".join(values))

            sheet_text = "\n".join(rows).strip()

            if sheet_text:
                sheets.append({
                    "page": sheet_name,
                    "text": sheet_text,
                    "source_type": "spreadsheet",
                })

    except Exception as exc:
        raise ValueError(
            f"Could not read spreadsheet: {exc}"
        ) from exc

    return sheets


# ============================================================
# CSV
# ============================================================

def extract_csv(file_path: str) -> list[dict]:
    """Extract CSV rows as readable text."""
    path = Path(file_path)

    try:
        dataframe = pd.read_csv(
            path,
            dtype=str,
            keep_default_na=False,
        )

        rows = []

        for _, row in dataframe.iterrows():

            values = [
                str(value).strip()
                for value in row.tolist()
                if str(value).strip()
            ]

            if values:
                rows.append(" | ".join(values))

        text = "\n".join(rows).strip()

        if not text:
            return []

        return [{
            "page": None,
            "text": text,
            "source_type": "csv",
        }]

    except Exception as exc:
        raise ValueError(
            f"Could not read CSV: {exc}"
        ) from exc


# ============================================================
# IMAGE / OCR
# ============================================================

def extract_image(file_path: str) -> list[dict]:
    """Extract text from an image using Tesseract OCR."""
    try:
        image = Image.open(file_path)

        text = pytesseract.image_to_string(
            image
        ).strip()

        if not text:
            return []

        return [{
            "page": None,
            "text": text,
            "source_type": "image_ocr",
        }]

    except Exception as exc:
        raise ValueError(
            f"Could not process image: {exc}"
        ) from exc


# ============================================================
# MAIN DISPATCHER
# ============================================================

def extract_text(file_path: str) -> list[dict]:
    """
    Automatically detect the file type and extract text.

    Supported:
        PDF
        TXT
        DOCX
        PPTX
        XLSX
        XLS
        CSV
        PNG
        JPG
        JPEG
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = path.suffix.lower()

    extractors = {
        ".pdf": extract_pdf,
        ".txt": extract_txt,
        ".docx": extract_docx,
        ".pptx": extract_pptx,
        ".xlsx": extract_excel,
        ".xls": extract_excel,
        ".csv": extract_csv,
        ".png": extract_image,
        ".jpg": extract_image,
        ".jpeg": extract_image,
    }

    extractor = extractors.get(extension)

    if extractor is None:
        supported = ", ".join(
            sorted(extractors.keys())
        )

        raise ValueError(
            f"Unsupported file type: {extension}. "
            f"Supported formats: {supported}"
        )

    return extractor(str(path))


# ============================================================
# CHUNKING
# ============================================================

def chunk_text(
    pages: list[dict],
    chunk_size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> list[dict]:
    """Split extracted text into overlapping chunks."""
    
    if chunk_size <= 0:
        raise ValueError(
            "chunk_size must be positive."
        )

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "overlap must be between 0 and chunk_size - 1."
        )

    chunks = []
    chunk_number = 0

    for page_data in pages:

        text = page_data["text"]
        start = 0

        while start < len(text):

            end = min(
                start + chunk_size,
                len(text)
            )

            # Prefer a nearby word boundary
            if end < len(text):

                boundary = text.rfind(
                    " ",
                    start + chunk_size // 2,
                    end,
                )

                if boundary > start:
                    end = boundary

            content = text[start:end].strip()

            if content:

                chunks.append({
                    "chunk_id": chunk_number,
                    "page": page_data.get("page"),
                    "text": content,
                    "start_char": start,
                    "end_char": end,
                    "source_type": page_data.get(
                        "source_type",
                        "unknown",
                    ),
                })

                chunk_number += 1

            if end >= len(text):
                break

            start = end - overlap

    return chunks